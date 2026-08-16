"""
Interior Design Project Management Workbook Builder - Python / Excel Version

Creates an .xlsx workbook with the same premium interior design project tracker
structure as the Google Sheets Apps Script version.

Requirements:
    pip install openpyxl

Run:
    python interior_design_project_workbook.py

Output:
    interior_design_project_workbook.xlsx

Notes:
- This is the Python Excel version. It creates a native .xlsx file.
- The Google Sheets version is still the best option if you want Google Apps
  Script automation inside Google Sheets.
- The Excel file can be uploaded to Google Sheets if needed.
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter, quote_sheetname
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as exc:  # fool-proof message for beginners
    raise SystemExit(
        "Missing dependency: openpyxl\n"
        "Install it with: pip install openpyxl\n"
        "Then run this script again."
    ) from exc

ROWS = 250
OUTPUT_FILE = "interior_design_project_workbook.xlsx"

COLORS = {
    "navy": "17365D",
    "blue": "1F4E78",
    "teal": "0F766E",
    "light_blue": "D9EAF7",
    "light_teal": "DDEFEA",
    "green": "D9EAD3",
    "yellow": "FFF2CC",
    "red": "F4CCCC",
    "gray": "D9D9D9",
    "white": "FFFFFF",
    "black": "000000",
}

LISTS: Dict[str, List[str]] = {
    "projectTypes": ["Residential", "Commercial", "Office", "Retail", "Hospitality", "Staging"],
    "yesNo": ["Yes", "No"],
    "projectStatuses": ["Not started", "In progress", "On hold", "Complete", "Delayed", "Cancelled"],
    "billingMethods": ["Flat fee", "Hourly", "Percentage", "Hybrid"],
    "designStatuses": ["Not started", "Concept", "Sourcing", "Client review", "Approved", "Ordered", "Installed", "Complete", "On hold"],
    "priorities": ["Low", "Medium", "High", "Urgent"],
    "roomTypes": ["Living room", "Bedroom", "Kitchen", "Bathroom", "Dining room", "Entryway", "Office", "Nursery", "Outdoor", "Commercial space", "Other"],
    "ffeCategories": ["Sofa", "Chair", "Table", "Bed", "Rug", "Lighting", "Artwork", "Mirror", "Window treatment", "Hardware", "Plumbing fixture", "Appliance", "Decor", "Textile", "Cabinetry", "Flooring", "Wallcovering", "Other"],
    "itemStatuses": ["Sourcing", "Option selected", "Client review", "Approved", "Ordered", "Shipped", "Delivered", "Installed", "Cancelled", "Backordered"],
    "orderStatuses": ["Not ordered", "Ordered", "Confirmed", "In production", "Shipped", "Delivered", "Delayed", "Cancelled", "Returned"],
    "approvalStatuses": ["Not sent", "Sent", "Client review", "Approved", "Revisions requested", "Rejected", "On hold"],
    "deliveryStatuses": ["Not ordered", "Ordered", "Shipped", "Delivered", "Delayed", "Cancelled", "Returned", "Backordered"],
    "installStatuses": ["Not scheduled", "Scheduled", "Installed", "Punch list", "Delayed", "Cancelled"],
    "fixtureTypes": ["Lighting", "Faucet", "Sink", "Shower fixture", "Tub", "Toilet", "Cabinet hardware", "Door hardware", "Appliance", "Electrical", "Other"],
    "finishTypes": ["Paint", "Wallpaper", "Flooring", "Tile", "Countertop", "Backsplash", "Cabinet finish", "Fabric", "Trim", "Stone", "Wood", "Metal", "Other"],
    "surfaces": ["Wall", "Ceiling", "Floor", "Cabinet", "Countertop", "Shower wall", "Backsplash", "Upholstery", "Drapery"],
    "vendorTypes": ["Furniture", "Lighting", "Fabric", "Flooring", "Tile", "Paint", "Wallpaper", "Decor", "Hardware", "Contractor", "Installer", "Workroom", "Custom furniture", "Other"],
    "budgetCategories": ["Furniture", "Lighting", "Rugs", "Decor", "Artwork", "Window treatments", "Flooring", "Paint/wallcovering", "Contractor labor", "Installation", "Shipping", "Taxes", "Design fee", "Contingency"],
    "damageChecks": ["Not checked", "Good condition", "Damaged", "Missing item", "Wrong item"],
    "approvalTypes": ["Furniture", "Finish", "Layout", "Color palette", "Budget", "Purchase", "Vendor", "Design concept"],
    "phases": ["Consultation", "Measurements", "Concept design", "Sourcing", "Client approval", "Purchasing", "Construction", "Delivery", "Installation", "Styling", "Final reveal", "Punch list"],
    "taskStatuses": ["Not started", "In progress", "Waiting on client", "Waiting on vendor", "Complete", "Delayed"],
    "measurementTypes": ["Room", "Wall", "Window", "Door", "Cabinet", "Furniture space", "Rug area", "Lighting location"],
    "trades": ["General contractor", "Painter", "Electrician", "Plumber", "Flooring installer", "Tile installer", "Wallpaper installer", "Carpenter", "Window treatment installer", "Furniture installer", "Cleaner", "Other"],
    "paymentStatuses": ["Draft", "Sent", "Paid", "Partially paid", "Overdue", "Cancelled", "Unpaid"],
    "meetingTypes": ["Client consultation", "Vendor meeting", "Site visit", "Installation meeting", "Contractor meeting", "Final walkthrough"],
    "preferredOptions": ["Option 1", "Option 2", "Option 3"],
}

SHEETS: List[Tuple[str, List[str]]] = [
    ("Welcome - Instructions", []),
    ("Project Dashboard", []),
    ("Client & Project Profile", []),
    ("Room List", ["Room name", "Room type", "Floor/level", "Square footage", "Design status", "Budget allocation", "Estimated spend", "Actual spend", "Remaining budget", "Priority level", "Notes"]),
    ("Master FF&E Schedule", ["Item ID", "Room", "Category", "Item name", "Description", "Quantity", "Dimensions", "Material", "Finish", "Color", "Vendor", "Brand", "Product link", "SKU/model number", "Unit cost", "Total cost", "Tax", "Shipping", "Discount", "Final cost", "Lead time", "Order status", "Approval status", "Delivery status", "Install status", "Notes", "Image link"]),
    ("Furniture Schedule", ["Room", "Furniture item", "Dimensions", "Quantity", "Vendor", "Product link", "Fabric/material", "Finish", "Cost", "Approval status", "Order status", "Delivery date", "Install date"]),
    ("Fixtures Schedule", ["Room", "Fixture type", "Item name", "Finish", "Size", "Vendor", "Model/SKU", "Quantity", "Unit cost", "Total cost", "Required rough-in date", "Order deadline", "Delivery date", "Installer notes"]),
    ("Finishes Schedule", ["Room", "Surface/application", "Finish type", "Product name", "Color/finish", "Material", "Vendor", "Product link", "SKU", "Quantity needed", "Unit cost", "Total cost", "Sample ordered", "Sample approved", "Installer notes"]),
    ("Product Sourcing Tracker", ["Room", "Item needed", "Option 1 link", "Option 1 price", "Option 2 link", "Option 2 price", "Option 3 link", "Option 3 price", "Preferred option", "Pros", "Cons", "Client feedback", "Final selection"]),
    ("Vendor & Supplier Database", ["Vendor name", "Vendor type", "Contact person", "Email", "Phone", "Website", "Account number", "Trade discount %", "Lead time", "Shipping terms", "Return policy", "Notes"]),
    ("Budget Tracker", ["Budget category", "Budgeted amount", "Estimated amount", "Actual amount", "Difference", "Paid amount", "Balance due", "Payment status", "Notes"]),
    ("Purchase Order Tracker", ["PO number", "Vendor", "Room", "Item", "Quantity", "Order date", "Expected ship date", "Expected delivery date", "Actual delivery date", "Order status", "Tracking number", "Carrier", "Invoice number", "Amount", "Paid status", "Notes"]),
    ("Delivery & Installation", ["Item", "Room", "Vendor", "Delivery address", "Expected delivery date", "Actual delivery date", "Delivery status", "Received by", "Damage check", "Damage notes", "Storage location", "Install date", "Installer", "Installation status", "Punch list notes"]),
    ("Client Approval Tracker", ["Item/design decision", "Room", "Approval type", "Sent to client date", "Client response date", "Status", "Approved option", "Client comments", "Revision needed", "Final approval date"]),
    ("Project Timeline", ["Phase", "Task", "Assigned to", "Start date", "Due date", "Completion date", "Status", "Priority", "Dependencies", "Notes", "Duration days", "Days overdue", "Gantt W1", "Gantt W2", "Gantt W3", "Gantt W4", "Gantt W5", "Gantt W6", "Gantt W7", "Gantt W8", "Gantt W9", "Gantt W10", "Gantt W11", "Gantt W12"]),
    ("Task & To-Do List", ["Task", "Room", "Category", "Assigned to", "Priority", "Due date", "Status", "Notes", "Days overdue"]),
    ("Measurement Log", ["Room", "Measurement type", "Width", "Height", "Depth", "Length", "Ceiling height", "Window size", "Door size", "Wall length", "Notes"]),
    ("Contractor Installer Tracker", ["Contractor/vendor name", "Trade", "Contact", "Scope", "Quote amount", "Start date", "End date", "Payment status", "Insurance received", "Contract signed", "Notes"]),
    ("Invoice Payment Tracker", ["Invoice number", "Invoice date", "Client", "Description", "Amount", "Due date", "Paid date", "Payment status", "Payment method", "Notes", "Balance due"]),
    ("Project Meeting Notes", ["Date", "Meeting type", "Attendees", "Discussion points", "Decisions made", "Action items", "Follow-up date", "Notes"]),
    ("Client FFE Report", ["Room", "Item name", "Description", "Quantity", "Vendor", "Product link", "Final cost", "Approval status", "Image link", "Client notes"]),
    ("Room Budget Report", ["Room", "Budget allocation", "Estimated spend", "Actual spend", "Remaining budget", "Budget used %", "Status"]),
    ("Vendor Order Report", ["Vendor", "PO number", "Room", "Item", "Quantity", "Order status", "Expected delivery date", "Amount", "Paid status", "Tracking number"]),
    ("Installation Checklist", ["Room", "Item", "Vendor", "Delivery status", "Install date", "Installer", "Installation status", "Damage check", "Punch list notes", "Complete?"]),
    ("Final Punch List", ["Room", "Issue/item", "Responsible party", "Priority", "Due date", "Status", "Resolution notes", "Completed date"]),
]


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def setup_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, _ in SHEETS:
        wb.create_sheet(sheet_name)
    create_lists_sheet(wb)
    setup_welcome(wb["Welcome - Instructions"])
    setup_profile(wb)
    for sheet_name, headers in SHEETS:
        if headers:
            setup_table_sheet(wb[sheet_name], headers)
    seed_budget_categories(wb)
    apply_validations(wb)
    apply_formulas(wb)
    setup_dashboard(wb)
    setup_printable_reports(wb)
    wb.active = wb.sheetnames.index("Welcome - Instructions")
    return wb


def create_lists_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Lists")
    for col_idx, (key, values) in enumerate(LISTS.items(), start=1):
        col = get_column_letter(col_idx)
        ws[f"{col}1"] = key
        ws[f"{col}1"].font = Font(bold=True)
        ws[f"{col}1"].fill = fill(COLORS["gray"])
        for row_idx, value in enumerate(values, start=2):
            ws[f"{col}{row_idx}"] = value
        ws.column_dimensions[col].width = max(14, len(key) + 2)
    ws.sheet_state = "hidden"


def setup_table_sheet(ws, headers: List[str]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ROWS + 1}"
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(1, col_idx, header)
        cell.font = Font(bold=True, color=COLORS["white"])
        cell.fill = fill(COLORS["blue"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = 18
    ws.row_dimensions[1].height = 42
    for row in range(2, ROWS + 2):
        if row % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row, col).fill = fill("F8FBFD")
    ws.protection.sheet = False
    for cell in ws[1]:
        cell.protection = Protection(locked=True)


def setup_welcome(ws) -> None:
    ws.sheet_view.showGridLines = False
    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 24
    ws.merge_cells("A1:F1")
    ws["A1"] = "Interior Design Project Management Workbook"
    ws["A1"].font = Font(size=20, bold=True, color=COLORS["white"])
    ws["A1"].fill = fill(COLORS["navy"])
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 48
    sections = [
        ("Quick Start Guide", "Step 1: Add client/project details in Client & Project Profile.\nStep 2: Add rooms in Room List.\nStep 3: Add FF&E items in Master FF&E Schedule.\nStep 4: Track budget, orders, approvals, delivery, and installation.\nStep 5: Use Project Dashboard for overview."),
        ("What Each Tab Is Used For", "Profile stores project details. Room List tracks rooms and budgets. Master FF&E is the core item schedule. Other tabs track furniture, fixtures, finishes, sourcing, vendors, budget, orders, delivery, approvals, timeline, tasks, measurements, contractors, invoices, and notes."),
        ("How To Enter Data", "Type into white cells beginning on row 2. Header rows and formula columns should not be overwritten unless you intentionally customize the workbook."),
        ("Color Coding", "Green = complete / approved / on budget. Yellow = pending / in progress. Red = overdue / over budget / delayed. Gray = cancelled / on hold."),
        ("Dropdown Instructions", "Click cells in status/category/room/vendor/priority columns and select from dropdown menus. Room and vendor dropdowns reference the Room List and Vendor Database."),
        ("Compatibility Note", "Built as a native Excel workbook with Python/openpyxl. It can be uploaded to Google Sheets. A Google Apps Script version is also included for native Google Sheets generation."),
        ("Editable / Unlocked Version Note", "This is editable. Header rows are styled to discourage accidental changes, but the workbook can be customized."),
        ("Support / Contact Note", "Add your support email, shop URL, license terms, or client handoff notes here before selling or sharing."),
    ]
    row = 3
    for title, text in sections:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row, 1, title).font = Font(bold=True, size=13, color=COLORS["white"])
        ws.cell(row, 1).fill = fill(COLORS["teal"])
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 2, end_column=6)
        ws.cell(row + 1, 1, text).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row + 1, 1).fill = fill("F8FBFD")
        row += 4


def setup_profile(wb: Workbook) -> None:
    ws = wb["Client & Project Profile"]
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "Client & Project Profile"
    ws["A1"].font = Font(size=18, bold=True, color=COLORS["white"])
    ws["A1"].fill = fill(COLORS["teal"])
    ws["A1"].alignment = Alignment(horizontal="center")
    fields = [
        ("Client name", ""), ("Client email", ""), ("Phone number", ""), ("Project name", ""), ("Project address", ""),
        ("Project type", "Residential"), ("Design style", ""), ("Start date", ""), ("Target completion date", ""), ("Total budget", 0),
        ("Designer name", ""), ("Client notes", ""), ("Scope summary", ""), ("Project status", "In progress"), ("Contract signed", "No"),
        ("Deposit paid", "No"), ("Design package selected", ""), ("Billing method", "Flat fee"),
    ]
    for idx, (label, value) in enumerate(fields, start=3):
        ws.cell(idx, 1, label).font = Font(bold=True)
        ws.cell(idx, 1).fill = fill(COLORS["light_teal"])
        ws.cell(idx, 2, value)
    ws["B12"].number_format = "$#,##0.00"
    ws["B10"].number_format = "m/d/yyyy"
    ws["B11"].number_format = "m/d/yyyy"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 35


def list_formula(list_key: str) -> str:
    col_idx = list(LISTS).index(list_key) + 1
    col = get_column_letter(col_idx)
    sheet = quote_sheetname("Lists")
    return f"={sheet}!${col}$2:${col}${len(LISTS[list_key]) + 1}"


def add_list_validation(ws, cell_range: str, list_key: str) -> None:
    dv = DataValidation(type="list", formula1=list_formula(list_key), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell_range)


def add_range_validation(ws, cell_range: str, formula: str) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell_range)


def apply_validations(wb: Workbook) -> None:
    room_formula = "='Room List'!$A$2:$A$251"
    vendor_formula = "='Vendor & Supplier Database'!$A$2:$A$251"
    profile = wb["Client & Project Profile"]
    for rng, key in [("B8", "projectTypes"), ("B17:B18", "yesNo"), ("B16", "projectStatuses"), ("B20", "billingMethods")]:
        add_list_validation(profile, rng, key)

    validation_map = {
        "Room List": {"B2:B251": "roomTypes", "E2:E251": "designStatuses", "J2:J251": "priorities"},
        "Master FF&E Schedule": {"C2:C251": "ffeCategories", "V2:V251": "orderStatuses", "W2:W251": "approvalStatuses", "X2:X251": "deliveryStatuses", "Y2:Y251": "installStatuses"},
        "Furniture Schedule": {"J2:J251": "approvalStatuses", "K2:K251": "orderStatuses"},
        "Fixtures Schedule": {"B2:B251": "fixtureTypes"},
        "Finishes Schedule": {"B2:B251": "surfaces", "C2:C251": "finishTypes", "M2:N251": "yesNo"},
        "Product Sourcing Tracker": {"I2:I251": "preferredOptions"},
        "Vendor & Supplier Database": {"B2:B251": "vendorTypes"},
        "Budget Tracker": {"A2:A251": "budgetCategories", "H2:H251": "paymentStatuses"},
        "Purchase Order Tracker": {"J2:J251": "orderStatuses", "O2:O251": "paymentStatuses"},
        "Delivery & Installation": {"G2:G251": "deliveryStatuses", "I2:I251": "damageChecks", "N2:N251": "installStatuses"},
        "Client Approval Tracker": {"C2:C251": "approvalTypes", "F2:F251": "approvalStatuses", "I2:I251": "yesNo"},
        "Project Timeline": {"A2:A251": "phases", "G2:G251": "taskStatuses", "H2:H251": "priorities"},
        "Task & To-Do List": {"E2:E251": "priorities", "G2:G251": "taskStatuses"},
        "Measurement Log": {"B2:B251": "measurementTypes"},
        "Contractor Installer Tracker": {"B2:B251": "trades", "H2:H251": "paymentStatuses", "I2:J251": "yesNo"},
        "Invoice Payment Tracker": {"H2:H251": "paymentStatuses"},
        "Project Meeting Notes": {"B2:B251": "meetingTypes"},
        "Final Punch List": {"D2:D251": "priorities", "F2:F251": "taskStatuses"},
    }
    for sheet_name, ranges in validation_map.items():
        for rng, key in ranges.items():
            add_list_validation(wb[sheet_name], rng, key)

    room_cells = [
        ("Master FF&E Schedule", "B2:B251"), ("Furniture Schedule", "A2:A251"), ("Fixtures Schedule", "A2:A251"),
        ("Finishes Schedule", "A2:A251"), ("Product Sourcing Tracker", "A2:A251"), ("Purchase Order Tracker", "C2:C251"),
        ("Delivery & Installation", "B2:B251"), ("Client Approval Tracker", "B2:B251"), ("Task & To-Do List", "B2:B251"),
        ("Measurement Log", "A2:A251"),
    ]
    vendor_cells = [
        ("Master FF&E Schedule", "K2:K251"), ("Furniture Schedule", "E2:E251"), ("Fixtures Schedule", "F2:F251"),
        ("Finishes Schedule", "G2:G251"), ("Purchase Order Tracker", "B2:B251"), ("Delivery & Installation", "C2:C251"),
    ]
    for sheet_name, rng in room_cells:
        add_range_validation(wb[sheet_name], rng, room_formula)
    for sheet_name, rng in vendor_cells:
        add_range_validation(wb[sheet_name], rng, vendor_formula)


def seed_budget_categories(wb: Workbook) -> None:
    ws = wb["Budget Tracker"]
    for idx, category in enumerate(LISTS["budgetCategories"], start=2):
        ws.cell(idx, 1, category)


def apply_formulas(wb: Workbook) -> None:
    for row in range(2, ROWS + 2):
        wb["Room List"][f"G{row}"] = f'=IFERROR(SUMIF(\'Master FF&E Schedule\'!$B:$B,$A{row},\'Master FF&E Schedule\'!$P:$P),0)'
        wb["Room List"][f"H{row}"] = f'=IFERROR(SUMIF(\'Master FF&E Schedule\'!$B:$B,$A{row},\'Master FF&E Schedule\'!$T:$T),0)'
        wb["Room List"][f"I{row}"] = f'=IF($A{row}="","",$F{row}-$H{row})'
        wb["Master FF&E Schedule"][f"P{row}"] = f'=IF($F{row}="","",$F{row}*$O{row})'
        wb["Master FF&E Schedule"][f"T{row}"] = f'=IF($P{row}="","",$P{row}+$Q{row}+$R{row}-$S{row})'
        wb["Fixtures Schedule"][f"J{row}"] = f'=IF($H{row}="","",$H{row}*$I{row})'
        wb["Finishes Schedule"][f"L{row}"] = f'=IF($J{row}="","",$J{row}*$K{row})'
        wb["Budget Tracker"][f"E{row}"] = f'=IF($A{row}="","",$B{row}-$D{row})'
        wb["Budget Tracker"][f"G{row}"] = f'=IF($A{row}="","",$D{row}-$F{row})'
        wb["Project Timeline"][f"K{row}"] = f'=IF(OR($D{row}="",$E{row}=""),"",$E{row}-$D{row}+1)'
        wb["Project Timeline"][f"L{row}"] = f'=IF(AND($E{row}<TODAY(),$G{row}<>"Complete"),TODAY()-$E{row},0)'
        for col in range(13, 25):
            week_offset = col - 13
            letter = get_column_letter(col)
            wb["Project Timeline"][f"{letter}{row}"] = f'=IF(OR($D{row}="",$E{row}=""),"",IF(AND($D{row}<=IF(\'Client & Project Profile\'!$B$10="",TODAY(),\'Client & Project Profile\'!$B$10)+{week_offset}*7+6,$E{row}>=IF(\'Client & Project Profile\'!$B$10="",TODAY(),\'Client & Project Profile\'!$B$10)+{week_offset}*7),"■",""))'
        wb["Task & To-Do List"][f"I{row}"] = f'=IF(AND($F{row}<TODAY(),$G{row}<>"Complete"),TODAY()-$F{row},0)'
        wb["Invoice Payment Tracker"][f"K{row}"] = f'=IF($A{row}="","",IF($H{row}="Paid",0,$E{row}))'
    add_budget_totals(wb["Budget Tracker"])
    apply_number_formats(wb)
    apply_conditional_formatting(wb)


def add_budget_totals(ws) -> None:
    start = ROWS + 4
    totals = [
        ("Total budget", f"=SUM(B2:B{ROWS+1})"),
        ("Total estimated", f"=SUM(C2:C{ROWS+1})"),
        ("Total actual", f"=SUM(D2:D{ROWS+1})"),
        ("Over/under budget", f"=B{start}-B{start+2}"),
        ("Total paid", f"=SUM(F2:F{ROWS+1})"),
        ("Remaining balance", f"=SUM(G2:G{ROWS+1})"),
    ]
    for idx, (label, formula) in enumerate(totals, start=start):
        ws.cell(idx, 1, label).font = Font(bold=True)
        ws.cell(idx, 1).fill = fill(COLORS["light_teal"])
        ws.cell(idx, 2, formula).number_format = "$#,##0.00"


def apply_number_formats(wb: Workbook) -> None:
    currency_ranges = {
        "Client & Project Profile": ["B12"],
        "Room List": ["F:I"],
        "Master FF&E Schedule": ["O:T"],
        "Furniture Schedule": ["I:I"],
        "Fixtures Schedule": ["I:J"],
        "Finishes Schedule": ["K:L"],
        "Product Sourcing Tracker": ["D:D", "F:F", "H:H"],
        "Budget Tracker": ["B:G"],
        "Purchase Order Tracker": ["N:N"],
        "Contractor Installer Tracker": ["E:E"],
        "Invoice Payment Tracker": ["E:E", "K:K"],
    }
    date_ranges = {
        "Furniture Schedule": ["L:M"], "Fixtures Schedule": ["K:M"], "Purchase Order Tracker": ["F:I"],
        "Delivery & Installation": ["E:F", "L:L"], "Client Approval Tracker": ["D:E", "J:J"],
        "Project Timeline": ["D:F"], "Task & To-Do List": ["F:F"], "Contractor Installer Tracker": ["F:G"],
        "Invoice Payment Tracker": ["B:B", "F:G"], "Project Meeting Notes": ["A:A", "G:G"], "Final Punch List": ["E:E", "H:H"],
    }
    for sheet, ranges in currency_ranges.items():
        for rng in ranges:
            set_range_format(wb[sheet], rng, "$#,##0.00")
    for sheet, ranges in date_ranges.items():
        for rng in ranges:
            set_range_format(wb[sheet], rng, "m/d/yyyy")


def set_range_format(ws, col_range: str, number_format: str) -> None:
    # Single cell, for example B12
    if ":" not in col_range:
        ws[col_range].number_format = number_format
        return

    # Explicit cell range, for example B2:G251
    if any(ch.isdigit() for ch in col_range):
        for row in ws[col_range]:
            cells = row if isinstance(row, tuple) else (row,)
            for cell in cells:
                cell.number_format = number_format
        return

    # Whole column or column block, for example B:G
    start_col, end_col = col_range.split(":")
    for col in range(ws[start_col + "1"].column, ws[end_col + "1"].column + 1):
        for row in range(2, ROWS + 2):
            ws.cell(row, col).number_format = number_format


def apply_conditional_formatting(wb: Workbook) -> None:
    green_words = ["Complete", "Approved", "Installed", "Paid"]
    yellow_words = ["Pending", "In progress", "Sent"]
    red_words = ["Delayed", "Overdue", "Rejected"]
    gray_words = ["Cancelled", "On hold"]
    for sheet_name, headers in SHEETS:
        if not headers:
            continue
        ws = wb[sheet_name]
        rng = f"A2:{get_column_letter(len(headers))}{ROWS + 1}"
        for word in green_words:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{word}",A2))'], fill=fill(COLORS["green"])))
        for word in yellow_words:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{word}",A2))'], fill=fill(COLORS["yellow"])))
        for word in red_words:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{word}",A2))'], fill=fill(COLORS["red"])))
        for word in gray_words:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{word}",A2))'], fill=fill(COLORS["gray"])))
    wb["Budget Tracker"].conditional_formatting.add(f"E2:E{ROWS+1}", CellIsRule(operator="lessThan", formula=["0"], fill=fill(COLORS["red"])))
    wb["Budget Tracker"].conditional_formatting.add(f"E2:E{ROWS+1}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=fill(COLORS["green"])))


def setup_dashboard(wb: Workbook) -> None:
    ws = wb["Project Dashboard"]
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:L1")
    ws["A1"] = "Project Dashboard"
    ws["A1"].font = Font(size=22, bold=True, color=COLORS["white"])
    ws["A1"].fill = fill(COLORS["navy"])
    ws["A1"].alignment = Alignment(horizontal="center")
    for col in range(1, 13):
        ws.column_dimensions[get_column_letter(col)].width = 16
    kpis = [
        ("Total project budget", "='Client & Project Profile'!B12"),
        ("Total estimated cost", "=SUM('Budget Tracker'!C2:C251)"),
        ("Total actual cost", "=SUM('Budget Tracker'!D2:D251)"),
        ("Remaining budget", "='Client & Project Profile'!B12-SUM('Budget Tracker'!D2:D251)"),
        ("Budget used %", "=IFERROR(SUM('Budget Tracker'!D2:D251)/'Client & Project Profile'!B12,0)"),
        ("Number of rooms", "=COUNTA('Room List'!A2:A251)"),
        ("Total FF&E items", "=COUNTA('Master FF&E Schedule'!D2:D251)"),
        ("Items approved", "=COUNTIF('Master FF&E Schedule'!W2:W251,\"Approved\")"),
        ("Items ordered", "=COUNTIF('Master FF&E Schedule'!V2:V251,\"Ordered\")+COUNTIF('Master FF&E Schedule'!V2:V251,\"Confirmed\")+COUNTIF('Master FF&E Schedule'!V2:V251,\"In production\")+COUNTIF('Master FF&E Schedule'!V2:V251,\"Shipped\")+COUNTIF('Master FF&E Schedule'!V2:V251,\"Delivered\")"),
        ("Items delivered", "=COUNTIF('Master FF&E Schedule'!X2:X251,\"Delivered\")"),
        ("Items installed", "=COUNTIF('Master FF&E Schedule'!Y2:Y251,\"Installed\")"),
        ("Overdue tasks", "=COUNTIF('Task & To-Do List'!I2:I251,\">0\")+COUNTIF('Project Timeline'!L2:L251,\">0\")"),
        ("Pending client approvals", "=COUNTIF('Client Approval Tracker'!F2:F251,\"Sent\")+COUNTIF('Client Approval Tracker'!F2:F251,\"Not sent\")+COUNTIF('Master FF&E Schedule'!W2:W251,\"Client review\")"),
        ("Open vendor orders", "=COUNTIF('Purchase Order Tracker'!J2:J251,\"Ordered\")+COUNTIF('Purchase Order Tracker'!J2:J251,\"Confirmed\")+COUNTIF('Purchase Order Tracker'!J2:J251,\"In production\")+COUNTIF('Purchase Order Tracker'!J2:J251,\"Shipped\")+COUNTIF('Purchase Order Tracker'!J2:J251,\"Delayed\")"),
        ("Items not ordered", "=COUNTIF('Master FF&E Schedule'!V2:V251,\"Not ordered\")+COUNTIFS('Master FF&E Schedule'!D2:D251,\"<>\",'Master FF&E Schedule'!V2:V251,\"\")"),
        ("Delivered not installed", "=COUNTIFS('Master FF&E Schedule'!X2:X251,\"Delivered\",'Master FF&E Schedule'!Y2:Y251,\"<>Installed\")"),
    ]
    for i, (label, formula) in enumerate(kpis):
        col = (i % 4) * 3 + 1
        row = 3 + (i // 4) * 3
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.cell(row, col, label).font = Font(bold=True, color=COLORS["white"])
        ws.cell(row, col).fill = fill(COLORS["blue"])
        ws.cell(row, col).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        ws.cell(row + 1, col, formula).font = Font(size=14, bold=True)
        ws.cell(row + 1, col).alignment = Alignment(horizontal="center")
        ws.cell(row + 1, col).fill = fill("F4F8FB")
    setup_dashboard_data_and_charts(ws)


def setup_dashboard_data_and_charts(ws) -> None:
    ws["A16"] = "Budget by Room"
    ws["A17"], ws["B17"], ws["C17"] = "Room", "Estimated", "Actual"
    for i in range(18, 31):
        src = i - 16
        ws[f"A{i}"] = f"='Room List'!A{src}"
        ws[f"B{i}"] = f"='Room List'!G{src}"
        ws[f"C{i}"] = f"='Room List'!H{src}"
    ws["E16"] = "Budget vs Actual"
    ws["E17"], ws["F17"] = "Metric", "Amount"
    ws["E18"], ws["F18"] = "Budget", "='Client & Project Profile'!B12"
    ws["E19"], ws["F19"] = "Actual", "=SUM('Budget Tracker'!D2:D251)"
    ws["H16"] = "Item Status Breakdown"
    ws["H17"], ws["I17"] = "Status", "Count"
    for idx, status in enumerate(LISTS["itemStatuses"], start=18):
        ws[f"H{idx}"] = status
        ws[f"I{idx}"] = f'=COUNTIF(\'Master FF&E Schedule\'!V2:Y251,"{status}")'
    ws["A35"] = "Delivery Status Breakdown"
    ws["A36"], ws["B36"] = "Status", "Count"
    for idx, status in enumerate(LISTS["deliveryStatuses"], start=37):
        ws[f"A{idx}"] = status
        ws[f"B{idx}"] = f'=COUNTIF(\'Delivery & Installation\'!G2:G251,"{status}")'
    ws["E35"] = "Category Spend Breakdown"
    ws["E36"], ws["F36"] = "Category", "Final Cost"
    for idx, category in enumerate(LISTS["ffeCategories"], start=37):
        ws[f"E{idx}"] = category
        ws[f"F{idx}"] = f'=SUMIF(\'Master FF&E Schedule\'!C2:C251,"{category}",\'Master FF&E Schedule\'!T2:T251)'

    chart1 = BarChart(); chart1.title = "Budget by Room"; chart1.add_data(Reference(ws, min_col=2, max_col=3, min_row=17, max_row=30), titles_from_data=True); chart1.set_categories(Reference(ws, min_col=1, min_row=18, max_row=30)); ws.add_chart(chart1, "A56")
    chart2 = BarChart(); chart2.title = "Budget vs Actual"; chart2.add_data(Reference(ws, min_col=6, min_row=17, max_row=19), titles_from_data=True); chart2.set_categories(Reference(ws, min_col=5, min_row=18, max_row=19)); ws.add_chart(chart2, "H56")
    chart3 = PieChart(); chart3.title = "Item Status Breakdown"; chart3.add_data(Reference(ws, min_col=9, min_row=18, max_row=27)); chart3.set_categories(Reference(ws, min_col=8, min_row=18, max_row=27)); ws.add_chart(chart3, "A72")
    chart4 = PieChart(); chart4.title = "Delivery Status Breakdown"; chart4.add_data(Reference(ws, min_col=2, min_row=37, max_row=44)); chart4.set_categories(Reference(ws, min_col=1, min_row=37, max_row=44)); ws.add_chart(chart4, "H72")
    chart5 = BarChart(); chart5.title = "Category Spend Breakdown"; chart5.add_data(Reference(ws, min_col=6, min_row=36, max_row=54), titles_from_data=True); chart5.set_categories(Reference(ws, min_col=5, min_row=37, max_row=54)); ws.add_chart(chart5, "A88")


def setup_printable_reports(wb: Workbook) -> None:
    for row in range(2, ROWS + 2):
        wb["Client FFE Report"][f"A{row}"] = f"='Master FF&E Schedule'!B{row}"
        wb["Client FFE Report"][f"B{row}"] = f"='Master FF&E Schedule'!D{row}"
        wb["Client FFE Report"][f"C{row}"] = f"='Master FF&E Schedule'!E{row}"
        wb["Client FFE Report"][f"D{row}"] = f"='Master FF&E Schedule'!F{row}"
        wb["Client FFE Report"][f"E{row}"] = f"='Master FF&E Schedule'!K{row}"
        wb["Client FFE Report"][f"F{row}"] = f"='Master FF&E Schedule'!M{row}"
        wb["Client FFE Report"][f"G{row}"] = f"='Master FF&E Schedule'!T{row}"
        wb["Client FFE Report"][f"H{row}"] = f"='Master FF&E Schedule'!W{row}"
        wb["Client FFE Report"][f"I{row}"] = f"='Master FF&E Schedule'!AA{row}"
        wb["Room Budget Report"][f"A{row}"] = f"='Room List'!A{row}"
        wb["Room Budget Report"][f"B{row}"] = f"='Room List'!F{row}"
        wb["Room Budget Report"][f"C{row}"] = f"='Room List'!G{row}"
        wb["Room Budget Report"][f"D{row}"] = f"='Room List'!H{row}"
        wb["Room Budget Report"][f"E{row}"] = f"='Room List'!I{row}"
        wb["Room Budget Report"][f"F{row}"] = f'=IFERROR(D{row}/B{row},0)'
        wb["Room Budget Report"][f"G{row}"] = f'=IF(E{row}>=0,"On budget","Over budget")'
        wb["Vendor Order Report"][f"A{row}"] = f"='Purchase Order Tracker'!B{row}"
        wb["Vendor Order Report"][f"B{row}"] = f"='Purchase Order Tracker'!A{row}"
        wb["Vendor Order Report"][f"C{row}"] = f"='Purchase Order Tracker'!C{row}"
        wb["Vendor Order Report"][f"D{row}"] = f"='Purchase Order Tracker'!D{row}"
        wb["Vendor Order Report"][f"E{row}"] = f"='Purchase Order Tracker'!E{row}"
        wb["Vendor Order Report"][f"F{row}"] = f"='Purchase Order Tracker'!J{row}"
        wb["Vendor Order Report"][f"G{row}"] = f"='Purchase Order Tracker'!H{row}"
        wb["Vendor Order Report"][f"H{row}"] = f"='Purchase Order Tracker'!N{row}"
        wb["Vendor Order Report"][f"I{row}"] = f"='Purchase Order Tracker'!O{row}"
        wb["Vendor Order Report"][f"J{row}"] = f"='Purchase Order Tracker'!K{row}"
        wb["Installation Checklist"][f"A{row}"] = f"='Delivery & Installation'!B{row}"
        wb["Installation Checklist"][f"B{row}"] = f"='Delivery & Installation'!A{row}"
        wb["Installation Checklist"][f"C{row}"] = f"='Delivery & Installation'!C{row}"
        wb["Installation Checklist"][f"D{row}"] = f"='Delivery & Installation'!G{row}"
        wb["Installation Checklist"][f"E{row}"] = f"='Delivery & Installation'!L{row}"
        wb["Installation Checklist"][f"F{row}"] = f"='Delivery & Installation'!M{row}"
        wb["Installation Checklist"][f"G{row}"] = f"='Delivery & Installation'!N{row}"
        wb["Installation Checklist"][f"H{row}"] = f"='Delivery & Installation'!I{row}"
        wb["Installation Checklist"][f"I{row}"] = f"='Delivery & Installation'!O{row}"
        wb["Installation Checklist"][f"J{row}"] = f'=IF(G{row}="Installed","Yes","No")'


def save_workbook(output_path: str | Path = OUTPUT_FILE) -> Path:
    wb = setup_workbook()
    path = Path(output_path).resolve()
    wb.save(path)
    return path


if __name__ == "__main__":
    saved = save_workbook()
    print(f"Workbook created successfully: {saved}")
